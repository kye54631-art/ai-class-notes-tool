# ============================
# 导出服务模块
# ============================
import os
import io
from datetime import datetime


class ExportService:
    """笔记导出服务 - 支持 TXT / Word / PDF"""

    @staticmethod
    def _sanitize_filename(title):
        """清理文件名中的非法字符"""
        import re
        title = re.sub(r'[\\/:*?"<>|]', '_', title)
        return title[:50]

    @staticmethod
    def export_txt(title, content):
        """导出为 TXT 纯文本"""
        full_text = f"=====================================\n"
        full_text += f"标题: {title}\n"
        full_text += f"导出时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
        full_text += f"=====================================\n\n"
        full_text += content
        return full_text.encode('utf-8'), f"{ExportService._sanitize_filename(title)}.txt", 'text/plain; charset=utf-8'

    @staticmethod
    def export_word(title, content):
        """导出为 Word (.docx)"""
        try:
            from docx import Document
            from docx.shared import Pt, Inches, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()

            # 设置默认字体
            style = doc.styles['Normal']
            style.font.name = '微软雅黑'
            style.font.size = Pt(12)

            # 标题
            title_para = doc.add_heading(title, level=0)
            title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # 导出时间
            time_para = doc.add_paragraph()
            time_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            time_run = time_para.add_run(f'导出时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
            time_run.font.size = Pt(9)
            time_run.font.color.rgb = RGBColor(128, 128, 128)

            doc.add_paragraph('─' * 50)

            # 将 Markdown 内容转为 Word 段落
            for line in content.split('\n'):
                line = line.strip()
                if not line:
                    doc.add_paragraph()
                    continue

                if line.startswith('## '):
                    doc.add_heading(line[3:], level=2)
                elif line.startswith('### '):
                    doc.add_heading(line[4:], level=3)
                elif line.startswith('#### '):
                    doc.add_heading(line[5:], level=4)
                elif line.startswith('- ') or line.startswith('* '):
                    doc.add_paragraph(line[2:], style='List Bullet')
                elif line.startswith('1. ') or line.startswith('2. '):
                    doc.add_paragraph(line[3:], style='List Number')
                elif line.startswith('> '):
                    para = doc.add_paragraph(line[2:])
                    para.paragraph_format.left_indent = Inches(0.5)
                else:
                    para = doc.add_paragraph(line)

            # 保存到内存
            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)

            filename = f"{ExportService._sanitize_filename(title)}.docx"
            mime = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            return buffer.getvalue(), filename, mime

        except ImportError:
            # 如果没有 python-docx，回退到 TXT
            return ExportService.export_txt(title, content)

    @staticmethod
    def export_pdf(title, content):
        """导出为 PDF"""
        try:
            from fpdf import FPDF

            pdf = FPDF()
            pdf.add_page()

            # 添加支持中文的字体
            # 在 Windows 上使用系统中文字体
            font_paths = [
                'C:/Windows/Fonts/msyh.ttc',   # 微软雅黑
                'C:/Windows/Fonts/simsun.ttc',  # 宋体
                'C:/Windows/Fonts/simhei.ttf',  # 黑体
            ]

            font_loaded = False
            for fp in font_paths:
                if os.path.exists(fp):
                    try:
                        pdf.add_font('CN', '', fp, uni=True)
                        pdf.add_font('CN', 'B', fp, uni=True)
                        font_loaded = True
                        break
                    except Exception:
                        continue

            if not font_loaded:
                # 无法加载中文字体，回退到 Word
                return ExportService.export_word(title, content)

            # 标题
            pdf.set_font('CN', 'B', 18)
            pdf.multi_cell(0, 12, title, align='C')
            pdf.ln(5)

            # 时间
            pdf.set_font('CN', '', 9)
            pdf.cell(0, 6, f'导出时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}', align='R')
            pdf.ln(8)

            # 分隔线
            pdf.set_draw_buffer(1)
            pdf.line(10, pdf.get_y(), 200, pdf.get_y())
            pdf.ln(5)

            # 内容
            pdf.set_font('CN', '', 11)
            for line in content.split('\n'):
                line = line.strip()
                if not line:
                    pdf.ln(3)
                    continue

                # 处理标题
                if line.startswith('## '):
                    pdf.set_font('CN', 'B', 14)
                    pdf.ln(3)
                    pdf.multi_cell(0, 8, line[3:])
                    pdf.set_font('CN', '', 11)
                elif line.startswith('### '):
                    pdf.set_font('CN', 'B', 12)
                    pdf.ln(2)
                    pdf.multi_cell(0, 7, line[4:])
                    pdf.set_font('CN', '', 11)
                elif line.startswith('#### '):
                    pdf.set_font('CN', 'B', 11)
                    pdf.multi_cell(0, 7, line[5:])
                    pdf.set_font('CN', '', 11)
                elif line.startswith('- ') or line.startswith('* '):
                    pdf.cell(5)
                    pdf.multi_cell(0, 6, '• ' + line[2:])
                elif line.startswith(('1. ', '2. ')):
                    pdf.multi_cell(0, 6, line)
                else:
                    # 移除 Markdown 粗体标记
                    clean = line.replace('**', '')
                    pdf.multi_cell(0, 6, clean)

            buffer = io.BytesIO()
            pdf.output(buffer)
            buffer.seek(0)

            filename = f"{ExportService._sanitize_filename(title)}.pdf"
            return buffer.getvalue(), filename, 'application/pdf'

        except ImportError:
            # 如果没有 fpdf2，回退到 Word
            return ExportService.export_word(title, content)

    @classmethod
    def export(cls, title, content, export_format='word'):
        """统一导出入口"""
        if export_format == 'txt':
            return cls.export_txt(title, content)
        elif export_format == 'word':
            return cls.export_word(title, content)
        elif export_format == 'pdf':
            return cls.export_pdf(title, content)
        else:
            return cls.export_txt(title, content)
